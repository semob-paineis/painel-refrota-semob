#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
 GERAR_PROPOSTAS.PY  —  Automação de Atualização do Radar de Propostas
 Ministério das Cidades / SEMOB — Painel REFROTA
================================================================================

O QUE ESTE SCRIPT FAZ
---------------------
Lê a aba "Lista de Projetos" da planilha "Dados_Refrota_Contratações.xlsx",
extrai a listagem completa de projetos (colunas B..J: Classificação, Tipo, UF,
Município, Proponente, Empreendimento, Apoio, Valor contratado, Situação) e
substitui o array `PROJETOS` embutido no arquivo "propostas.html", mantendo
todo o restante do template (filtros, estilos, lógica de tabela) intacto.

Passa a ser parte do pipeline semanal, junto com gerar_dados.py:

    python gerar_dados.py
    python gerar_propostas.py

COMO USAR
---------
    python gerar_propostas.py
    python gerar_propostas.py --planilha "caminho/para/arquivo.xlsx" \
                               --template propostas.html \
                               --saida propostas.html

REQUISITOS
----------
    pip install openpyxl
================================================================================
"""

import argparse
import json
import re
import sys
from datetime import datetime
try:
    from zoneinfo import ZoneInfo
    FUSO_BRASILIA = ZoneInfo("America/Sao_Paulo")
except ImportError:
    from datetime import timezone, timedelta
    FUSO_BRASILIA = timezone(timedelta(hours=-3))

import openpyxl

ABA_LISTA = "Lista de Projetos"
LINHA_CABECALHO = 3       # linha 3: "Classificação", "Tipo", "UF", ...
LINHA_PRIMEIRA_DADO = 4   # dados começam na linha 4

# Colunas fixas (B..J), 1-indexadas (B=2)
COL_CLASSIFICACAO = 2
COL_TIPO = 3
COL_UF = 4
COL_MUNICIPIO = 5
COL_PROPONENTE = 6
COL_EMPREENDIMENTO = 7
COL_APOIO = 8
COL_VALOR_CONTRATADO = 9
COL_SITUACAO = 10


def log(nivel, msg):
    prefixo = {"ok": "[OK]   ", "info": "[INFO] ", "erro": "[ERRO] ", "aviso": "[AVISO]"}
    print(f"  {prefixo.get(nivel, '')} {msg}")


def limpar_texto(v):
    """Remove espaços supérfluos que existem na planilha (ex.: 'Em ação
    preparatória ' com espaço à direita), preservando None."""
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def extrair_projetos(caminho_planilha):
    wb = openpyxl.load_workbook(caminho_planilha, data_only=True)
    if ABA_LISTA not in wb.sheetnames:
        raise SystemExit(f"Aba '{ABA_LISTA}' não encontrada na planilha.")
    ws = wb[ABA_LISTA]

    # Confirma cabeçalho esperado (alerta, não interrompe, caso a planilha mude)
    esperado = {
        COL_CLASSIFICACAO: "Classificação", COL_TIPO: "Tipo", COL_UF: "UF",
        COL_MUNICIPIO: "Município Principal", COL_PROPONENTE: "Proponente",
        COL_EMPREENDIMENTO: "Empreendimento", COL_APOIO: "Apoio",
        COL_VALOR_CONTRATADO: "Valor contratado", COL_SITUACAO: "Situação Contrato",
    }
    for col, titulo in esperado.items():
        real = ws.cell(LINHA_CABECALHO, col).value
        real_norm = (real or "").strip()
        if real_norm != titulo:
            log("aviso", f"Cabeçalho da coluna {col} é '{real_norm}', esperado '{titulo}' "
                          f"— confira se a planilha mudou de layout.")

    # Encontra onde a listagem termina: primeira linha "TOTAL" na coluna B,
    # ou fim da planilha. Linhas em branco no meio são puladas (não encerram).
    linha_fim = ws.max_row
    for r in range(LINHA_PRIMEIRA_DADO, ws.max_row + 1):
        b = ws.cell(r, COL_CLASSIFICACAO).value
        if isinstance(b, str) and b.strip().upper().startswith("TOTAL"):
            linha_fim = r - 1
            break

    projetos = []
    for r in range(LINHA_PRIMEIRA_DADO, linha_fim + 1):
        classificacao = limpar_texto(ws.cell(r, COL_CLASSIFICACAO).value)
        if not classificacao:
            continue  # linha em branco entre projetos — ignora
        projetos.append({
            "classificacao": classificacao,
            "tipo": limpar_texto(ws.cell(r, COL_TIPO).value),
            "uf": limpar_texto(ws.cell(r, COL_UF).value),
            "municipio": limpar_texto(ws.cell(r, COL_MUNICIPIO).value),
            "proponente": limpar_texto(ws.cell(r, COL_PROPONENTE).value),
            "empreendimento": limpar_texto(ws.cell(r, COL_EMPREENDIMENTO).value),
            "apoio": ws.cell(r, COL_APOIO).value,
            "valorContratado": ws.cell(r, COL_VALOR_CONTRATADO).value,
            "situacao": limpar_texto(ws.cell(r, COL_SITUACAO).value),
        })

    return projetos


def comparar_com_anterior(novos, caminho_template):
    """Lê o PROJETOS embutido no template atual e reporta o que mudou,
    para conferência (mesmo espírito do comparativo do gerar_dados.py)."""
    try:
        html = open(caminho_template, encoding="utf-8").read()
        m = re.search(r"const PROJETOS = (\[.*?\]);", html, re.DOTALL)
        if not m:
            return
        antigos = json.loads(m.group(1))
    except Exception:
        return

    def chave(p):
        return (p.get("uf"), p.get("municipio"), p.get("proponente"),
                p.get("empreendimento"), p.get("apoio"))

    mapa_antigo = {chave(p): p for p in antigos}
    mapa_novo = {chave(p): p for p in novos}

    novos_ids = [p for k, p in mapa_novo.items() if k not in mapa_antigo]
    removidos_ids = [p for k, p in mapa_antigo.items() if k not in mapa_novo]
    mudancas = []
    for k, novo in mapa_novo.items():
        antigo = mapa_antigo.get(k)
        if antigo and (antigo.get("situacao") != novo.get("situacao")
                       or antigo.get("valorContratado") != novo.get("valorContratado")):
            mudancas.append((antigo, novo))

    log("info", f"Total antes: {len(antigos)} | Total agora: {len(novos)}")
    if novos_ids:
        log("info", f"{len(novos_ids)} projeto(s) novo(s) na listagem.")
    if removidos_ids:
        log("info", f"{len(removidos_ids)} projeto(s) não encontrados mais (verifique se não "
                     f"foi só mudança de proponente/UF/valor usada como chave).")
    if mudancas:
        log("info", f"{len(mudancas)} projeto(s) com mudança real de situação/valor contratado:")
        for antigo, novo in mudancas:
            print(f"           {novo.get('uf')} — {novo.get('municipio')} — "
                  f"{novo.get('empreendimento')}: "
                  f"'{antigo.get('situacao')}' -> '{novo.get('situacao')}'"
                  + ("" if antigo.get('valorContratado') == novo.get('valorContratado')
                     else f" | valor {antigo.get('valorContratado')} -> {novo.get('valorContratado')}"))


def injetar_no_template(projetos, caminho_template, caminho_saida, data_str):
    html = open(caminho_template, encoding="utf-8").read()

    projetos_js = json.dumps(projetos, ensure_ascii=False, separators=(",", ":"))
    novo_html, n = re.subn(
        r"const PROJETOS = \[.*?\];",
        "const PROJETOS = " + projetos_js.replace("\\", "\\\\") + ";",
        html, count=1, flags=re.DOTALL
    )
    if n == 0:
        raise SystemExit("Não encontrei 'const PROJETOS = [...]' no template — "
                          "verifique se o arquivo é o propostas.html correto.")

    novo_html, n2 = re.subn(
        r"const DATA_ATUALIZACAO = '[^']*';",
        f"const DATA_ATUALIZACAO = '{data_str}';",
        novo_html, count=1
    )
    if n2 == 0:
        log("aviso", "Não encontrei a constante DATA_ATUALIZACAO para atualizar (não é bloqueante).")

    with open(caminho_saida, "w", encoding="utf-8") as f:
        f.write(novo_html)


def main():
    parser = argparse.ArgumentParser(description="Gera propostas.html a partir da aba 'Lista de Projetos'.")
    parser.add_argument("--planilha", default="Dados_Refrota_Contratações.xlsx")
    parser.add_argument("--template", default="propostas.html",
                         help="Arquivo propostas.html a ser usado como base (estrutura/estilos).")
    parser.add_argument("--saida", default="propostas.html")
    args = parser.parse_args()

    agora = datetime.now(FUSO_BRASILIA)
    print("=" * 72)
    print(" ATUALIZAÇÃO DO RADAR DE PROPOSTAS — geração de propostas.html")
    print(f" {agora.strftime('%d/%m/%Y %H:%M')}")
    print("=" * 72)

    log("info", f"Lendo aba '{ABA_LISTA}' de {args.planilha}")
    projetos = extrair_projetos(args.planilha)
    log("ok", f"{len(projetos)} projetos extraídos.")

    print("-" * 72)
    comparar_com_anterior(projetos, args.template)

    print("-" * 72)
    data_str = agora.strftime("%d/%m/%Y")
    injetar_no_template(projetos, args.template, args.saida, data_str)
    log("ok", f"Arquivo gerado: {args.saida} (data de atualização: {data_str})")
    print("=" * 72)


if __name__ == "__main__":
    main()
