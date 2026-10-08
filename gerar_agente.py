#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gerar_agente.py — gera base_agente.js, a base por contrato usada pelo
Agente de Dúvidas do painel REFROTA.

Lê a aba BASEDEDADOS (uma linha por contrato/proposta) da planilha semanal e
grava `base_agente.js` (window.BASE_AGENTE = {...}). O agente roda inteiro no
navegador e calcula toda resposta a partir dessas linhas.

Uso:
    python3 gerar_agente.py --planilha Dados_Refrota_Contratações.xlsx
    python3 gerar_agente.py --planilha ... --dados dados.json --saida base_agente.js

Se --dados for informado (padrão: dados.json ao lado do script), a base é
CONFERIDA contra os totais do painel (propostas, veículos, investimento,
elétricos). Qualquer divergência é exibida e o script termina com erro, para
que o agente nunca responda com um número diferente do que o painel mostra.

ATENÇÃO (mesma regra do gerar_propostas.py): rode contra a cópia ORIGINAL da
planilha, nunca contra uma cópia salva pelo openpyxl (que apaga o cache das
fórmulas). BASEDEDADOS é lida com data_only=True.
"""
import argparse
import json
import os
import sys
from datetime import datetime, date

import openpyxl

# Colunas da BASEDEDADOS (cabeçalho na linha 1). Índice 0-based.
COLUNAS = [
    # chave,       índice, rótulo na planilha
    ("proposta",     0, "Proposta"),
    ("anoProposta",  1, "Ano Proposta"),
    ("contrato",     4, "Contrato"),
    ("tipo",         5, "Tipo"),
    ("uf",           6, "UF"),
    ("regiao",       7, "Região"),
    ("proponente",   8, "Proponente"),
    ("tipoProp",     9, "Tipo de Proponente"),
    ("municipio",   10, "Município Principal"),
    ("agente",      11, "Agente Financeiro"),
    ("programa",    12, "Programa"),
    ("empreend",    13, "Empreendimento"),
    ("fonte",       14, "Fonte"),
    ("apoio",       15, "Apoio (Previsto Novo PAC)"),
    ("contrapartida", 17, "Contrapartida"),
    ("situacao",    18, "Situação Contrato"),
    ("assinatura",  19, "Data assinatura"),
    ("execucao",    20, "Situação da execução"),
    ("valorContr",  21, "Valor contratado"),
    ("liberado",    22, "Valor liberado/desembolsado"),
    ("qtdContr",    23, "Qtd. contratada"),
    ("elContr",     24, "Elétricos Contratados"),
    ("e6Contr",     25, "Euro 6 Contratados"),
    ("trContr",     26, "Trilhos Contratados"),
    ("qtdEntregue", 27, "Qtd. entregue"),
    ("portaria",    29, "Portaria"),
    ("anoPortaria", 30, "Ano portaria"),
    ("item",        31, "Item principal"),
    ("qtdSel",      33, "Qnt (selecionada)"),
    ("elSel",       34, "Elétricos Selecionados"),
    ("e6Sel",       35, "Euro 6 Selecionados"),
    ("trSel",       36, "Trilhos Selecionados"),
    ("fisico",      38, "% físico"),
    ("fin",         39, "% fin"),
]


def categoria(situacao):
    """Mesma regra de classificação da aba 'Dados para Painel'."""
    s = (situacao or "").lower()
    if "habilit" in s:
        return "habilitada"          # fora da carteira selecionada
    if "cancel" in s:
        return "cancelada"
    if s.strip() == "desistência":
        return "acancelar"
    if "prepa" in s:
        return "preparacao"
    if any(k in s for k in ("contrat", "concl", "andamento", "licita")):
        return "contratada"
    return "outra"


def limpar(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return round(v, 4) if v != int(v) else int(v)
    if isinstance(v, str):
        return v.strip()
    return v


def ler_linhas(planilha):
    wb = openpyxl.load_workbook(planilha, data_only=True, read_only=True)
    if "BASEDEDADOS" not in wb.sheetnames:
        sys.exit("Aba 'BASEDEDADOS' não encontrada na planilha.")
    ws = wb["BASEDEDADOS"]
    linhas = []
    for numero, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue
        row = list(row) + [None] * 45
        reg = [limpar(row[i]) for _, i, _ in COLUNAS]
        reg.append(categoria(row[18]))
        reg.append(numero)   # linha na planilha (rastreabilidade)
        linhas.append(reg)
    return linhas


def conferir(linhas, caminho_dados):
    """Confere a base contra os totais do painel. Devolve lista de divergências."""
    if not caminho_dados or not os.path.exists(caminho_dados):
        return ["dados.json não encontrado — conferência pulada."], False
    d = json.load(open(caminho_dados, encoding="utf-8"))["consolidado"]
    ix = {c[0]: n for n, c in enumerate(COLUNAS)}
    icat = len(COLUNAS)
    n = lambda r, k: r[ix[k]] or 0
    refrota = [r for r in linhas if "refrota" in (r[ix["tipo"]] or "").lower()]
    contr = [r for r in refrota if r[icat] == "contratada"]
    sel = [r for r in refrota if r[icat] != "habilitada"]
    esperado = {
        "propostas contratadas": (len(contr), d["projetos"]["contratados"]["propostas"]),
        "veículos contratados": (sum(n(r, "qtdContr") for r in contr),
                                 sum(d["veiculos"].values())),
        "elétricos contratados": (sum(n(r, "elContr") for r in contr), d["veiculos"]["eletricos"]),
        "euro 6 contratados": (sum(n(r, "e6Contr") for r in contr), d["veiculos"]["euro6"]),
        "investimento contratado": (round(sum(n(r, "valorContr") for r in contr), 2),
                                    round(d["projetos"]["contratados"]["investimento"], 2)),
        "propostas selecionadas": (len(sel), d["projetos"]["selecionados"]["propostas"]),
        "veículos selecionados": (sum(n(r, "elSel") + n(r, "e6Sel") + n(r, "trSel") for r in sel),
                                  sum(d["veiculosSelecionados"].values())),
        "investimento selecionado": (round(sum(n(r, "apoio") for r in sel), 2),
                                     round(d["projetos"]["selecionados"]["investimento"], 2)),
        "em preparação": (sum(1 for r in refrota if r[icat] == "preparacao"),
                          d["projetos"]["status"]["emPreparacao"]["qtd"]),
    }
    erros = [f"{k}: base={a} ≠ painel={b}" for k, (a, b) in esperado.items() if abs(a - b) > 0.5]
    return erros, True


def entregues_por_ano(linhas, colunas):
    """
    Veículos ENTREGUES por ano da portaria de seleção, por cenário (consolidado/publico/privado).
    A planilha não registra a data de cada entrega, então o ano é o da portaria da proposta
    (o mesmo "Ano" usado nas tabelas de Evolução por Ano do painel).
    """
    ix = {c: i for i, c in enumerate(colunas)}
    out = {"consolidado": {}, "publico": {}, "privado": {}}
    for r in linhas:
        ent = r[ix["qtdEntregue"]] or 0
        ano = r[ix["anoPortaria"]]
        if not ent or ano is None:
            continue
        ano = str(int(ano))
        frente = "privado" if r[ix["tipo"]] == "Refrota Privado" else "publico"
        for cen in ("consolidado", frente):
            out[cen][ano] = out[cen].get(ano, 0) + ent
    return {cen: dict(sorted(v.items())) for cen, v in out.items()}


def main():
    ap = argparse.ArgumentParser(description="Gera base_agente.js (base por contrato do Agente de Dúvidas).")
    ap.add_argument("--planilha", default="Dados_Refrota_Contratações.xlsx")
    ap.add_argument("--dados", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "dados.json"),
                    help="dados.json do painel, usado para conferir os totais.")
    ap.add_argument("--saida", default="base_agente.js")
    ap.add_argument("--forcar", action="store_true", help="Grava mesmo com divergências (só para inspeção).")
    a = ap.parse_args()

    if not os.path.exists(a.planilha):
        sys.exit(f"Planilha não encontrada: {a.planilha}")
    linhas = ler_linhas(a.planilha)
    print(f"  [OK]    {len(linhas)} contratos lidos de BASEDEDADOS.")

    erros, conferido = conferir(linhas, a.dados)
    if erros:
        for e in erros:
            print(f"  [ERRO]  {e}" if conferido else f"  [AVISO] {e}")
        if conferido and not a.forcar:
            sys.exit("  [ERRO]  A base não fecha com o painel — arquivo NÃO gerado (use --forcar só para inspeção).")
    else:
        print("  [OK]    Base conferida: fecha com os totais do painel (propostas, veículos, elétricos, Euro 6, investimentos).")

    meta = None
    if os.path.exists(a.dados):
        meta = (json.load(open(a.dados, encoding="utf-8")).get("_meta") or {}).get("gerado_em")
    base = {
        "geradoEm": meta,
        "planilha": os.path.basename(a.planilha),
        "colunas": [c[0] for c in COLUNAS] + ["cat", "linha"],
        "linhas": linhas,
        "entreguesPorAno": entregues_por_ano(linhas, [c[0] for c in COLUNAS] + ["cat", "linha"]),
    }
    with open(a.saida, "w", encoding="utf-8") as f:
        f.write("/* Gerado por gerar_agente.py — não editar à mão. */\n")
        f.write("window.BASE_AGENTE = ")
        json.dump(base, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
    print(f"  [OK]    Arquivo gerado: {a.saida} ({os.path.getsize(a.saida)/1024:.0f} KB)")


if __name__ == "__main__":
    main()
